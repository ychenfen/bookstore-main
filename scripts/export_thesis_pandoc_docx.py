#!/usr/bin/env python3
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
THESIS_DIR = ROOT / "thesis_assets"
TEX_PATH = THESIS_DIR / "thesis.tex"
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_DOCX = OUTPUT_DIR / "bookstore_thesis.docx"
OUTPUT_PDF = OUTPUT_DIR / "bookstore_thesis.pdf"


def find_pandoc() -> Path:
    path = shutil.which("pandoc")
    if path:
        return Path(path)

    candidates = sorted((ROOT / "tmp" / "pandoc").glob("pandoc-*/bin/pandoc"))
    if candidates:
        return candidates[-1]

    raise SystemExit("pandoc not found. Install pandoc or place it under tmp/pandoc/")


def set_style_font(style, western: str, east_asia: str, size_pt: int, *, bold: bool = False) -> None:
    style.font.name = western
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), western)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), western)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_run_font(run, western: str, east_asia: str, size_pt: int, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)
    rfonts.set(qn("w:eastAsia"), east_asia)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def add_bottom_border(paragraph) -> None:
    ppr = paragraph._element.get_or_add_pPr()
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        ppr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    p_bdr.append(bottom)


def build_reference_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(30)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(30)
    section.right_margin = Mm(25)

    normal = doc.styles["Normal"]
    set_style_font(normal, "Times New Roman", "SimSun", 12)

    heading1 = doc.styles["Heading 1"]
    set_style_font(heading1, "Arial", "SimHei", 16, bold=True)
    heading1.paragraph_format.space_before = Pt(18)
    heading1.paragraph_format.space_after = Pt(12)

    heading2 = doc.styles["Heading 2"]
    set_style_font(heading2, "Arial", "SimHei", 14, bold=True)
    heading2.paragraph_format.space_before = Pt(14)
    heading2.paragraph_format.space_after = Pt(8)

    heading3 = doc.styles["Heading 3"]
    set_style_font(heading3, "Arial", "SimHei", 12, bold=True)
    heading3.paragraph_format.space_before = Pt(10)
    heading3.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    set_style_font(title, "Arial", "SimHei", 18, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)

    doc.add_paragraph("reference")
    doc.save(path)


def clean_tex(raw: str) -> str:
    body_match = re.search(r"\\begin\{document\}(.*)\\end\{document\}", raw, re.S)
    body = body_match.group(1) if body_match else raw

    body = re.sub(r"^\s*\\thispagestyle\{[^}]+\}\s*$", "", body, flags=re.M)
    body = re.sub(r"^\s*\\pagestyle\{[^}]+\}\s*$", "", body, flags=re.M)
    body = re.sub(r"^\s*\\pagenumbering\{[^}]+\}\s*$", "", body, flags=re.M)
    body = re.sub(r"^\s*\\setcounter\{page\}\{[^}]+\}\s*$", "", body, flags=re.M)
    body = re.sub(r"^\s*\\addcontentsline\{[^}]+\}\{[^}]+\}\{[^}]+\}\s*$", "", body, flags=re.M)
    body = re.sub(r"^\s*\\tableofcontents\s*$", "", body, flags=re.M)

    def strip_equation_label(match: re.Match[str]) -> str:
        block = match.group(0)
        return re.sub(r"\s*\\label\{[^}]+\}", "", block)

    body = re.sub(
        r"\\begin\{equation\}.*?\\end\{equation\}",
        strip_equation_label,
        body,
        flags=re.S,
    )

    body = re.sub(r"\n{3,}", "\n\n", body)
    return body.strip() + "\n"


def run_pandoc(clean_tex_path: Path, reference_docx: Path, output_docx: Path) -> None:
    pandoc = find_pandoc()
    command = [
        str(pandoc),
        str(clean_tex_path),
        "--from=latex",
        "--to=docx",
        "--resource-path",
        str(THESIS_DIR),
        "--number-sections",
        "--reference-doc",
        str(reference_docx),
        "-o",
        str(output_docx),
    ]
    subprocess.run(command, check=True, cwd=THESIS_DIR)


def add_page_break_before(paragraph) -> None:
    paragraph.paragraph_format.page_break_before = True


def rebuild_cover(doc: Document) -> None:
    abstract_paragraph = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "摘要":
            abstract_paragraph = paragraph
            break

    if abstract_paragraph is None:
        return

    for paragraph in list(doc.paragraphs):
        if paragraph == abstract_paragraph:
            break
        remove_paragraph(paragraph)

    cover_lines = [
        ("江西师大软件学院", 16, True, False, 0),
        ("本科生毕业论文", 16, True, False, 0),
        ("", 10, False, False, 0),
        ("基于React与tRPC的网上书城系统", 18, True, False, 0),
        ("设计与实现", 18, True, False, 0),
        ("", 10, False, False, 0),
        ("Design and Implementation of an Online Bookstore System", 12, False, True, 0),
        ("Based on React and tRPC", 12, False, True, 0),
        ("", 10, False, False, 0),
        ("", 10, False, False, 0),
        ("学生姓名：", 12, False, False, 0),
        ("", 10, False, False, 0),
        ("学号：", 12, False, False, 0),
        ("", 10, False, False, 0),
        ("所在学院： 软件学院", 12, False, False, 0),
        ("", 10, False, False, 0),
        ("所学专业： 软件工程", 12, False, False, 0),
        ("", 10, False, False, 0),
        ("指导教师：", 12, False, False, 0),
        ("", 10, False, False, 0),
        ("完成时间：", 12, False, False, 0),
    ]

    for text, size, bold, italic, after in reversed(cover_lines):
        paragraph = abstract_paragraph.insert_paragraph_before(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(after)
        run = paragraph.runs[0] if paragraph.runs else paragraph.add_run("")
        set_run_font(
            run,
            "Times New Roman" if italic else "Arial",
            "SimHei" if bold else "SimSun",
            size,
            bold=bold,
            italic=italic,
        )

    add_page_break_before(abstract_paragraph)


def apply_header(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.text = "基于 React 与 tRPC 的网上书城系统设计与实现"
    if paragraph.runs:
        set_run_font(paragraph.runs[0], "Times New Roman", "SimSun", 8)
    add_bottom_border(paragraph)


def insert_toc(doc: Document) -> None:
    headings: list[tuple[int, str]] = []
    first_heading = None

    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        text = paragraph.text.replace("\t", " ").strip()
        if not text:
            continue
        if style_name == "Heading 1":
            headings.append((1, text))
            if first_heading is None:
                first_heading = paragraph
        elif style_name == "Heading 2":
            headings.append((2, text))

    if first_heading is None or not headings:
        return

    toc_title = first_heading.insert_paragraph_before("目录")
    toc_title.style = doc.styles["Heading 1"]
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_after = Pt(12)
    toc_title.paragraph_format.page_break_before = True

    for level, text in headings:
        paragraph = first_heading.insert_paragraph_before(text)
        paragraph.style = doc.styles["Normal"]
        paragraph.paragraph_format.left_indent = Mm(0 if level == 1 else 8)
        paragraph.paragraph_format.space_after = Pt(3)
    first_heading.paragraph_format.page_break_before = True


def polish_docx(path: Path) -> None:
    doc = Document(path)

    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(30)
        section.bottom_margin = Mm(25)
        section.left_margin = Mm(30)
        section.right_margin = Mm(25)
        apply_header(section)

    rebuild_cover(doc)

    non_empty = [p for p in doc.paragraphs if p.text.strip()]
    if len(non_empty) >= 2:
        for paragraph in non_empty:
            if paragraph.text.strip() == "摘要":
                paragraph.style = doc.styles["Heading 1"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif paragraph.text.strip() == "Abstract":
                paragraph.style = doc.styles["Heading 1"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    insert_toc(doc)

    heading_break_done = False
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""

        if text in {"摘 要", "Abstract"}:
            paragraph.paragraph_format.page_break_before = True
            continue

        if style_name == "Heading 1" and text:
            if heading_break_done:
                paragraph.paragraph_format.page_break_before = True
            heading_break_done = True

    doc.save(path)


def render_pdf(docx_path: Path, pdf_path: Path) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [
            "soffice",
            "-env:UserInstallation=file:///tmp/lo_profile_pandoc_export",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(pdf_path.parent),
            str(docx_path),
        ],
        check=True,
        cwd=ROOT,
    )


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    raw = TEX_PATH.read_text(encoding="utf-8")

    with tempfile.TemporaryDirectory(prefix="thesis-pandoc-") as temp_dir_str:
        temp_dir = Path(temp_dir_str)
        clean_tex_path = temp_dir / "thesis_clean.tex"
        reference_docx = temp_dir / "reference.docx"

        clean_tex_path.write_text(clean_tex(raw), encoding="utf-8")
        build_reference_docx(reference_docx)
        run_pandoc(clean_tex_path, reference_docx, OUTPUT_DOCX)

    polish_docx(OUTPUT_DOCX)
    render_pdf(OUTPUT_DOCX, OUTPUT_PDF)


if __name__ == "__main__":
    main()
