#!/usr/bin/env python3

import re
from pathlib import Path
from typing import Dict, List, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
THESIS_DIR = ROOT / "thesis_assets"
TEX_PATH = THESIS_DIR / "thesis.tex"
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_PATH = OUTPUT_DIR / "bookstore_thesis.docx"


def set_east_asia_font(style, east_asia: str, latin: str = "Times New Roman") -> None:
    style.font.name = latin
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size in [("Title", 18), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
      style = doc.styles[style_name]
      style.font.bold = True
      style.font.size = Pt(size)
      style.font.name = "Times New Roman"
      style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    if "CaptionCn" not in [s.name for s in doc.styles]:
        caption_style = doc.styles.add_style("CaptionCn", WD_STYLE_TYPE.PARAGRAPH)
        set_east_asia_font(caption_style, "宋体")
        caption_style.paragraph_format.space_before = Pt(6)
        caption_style.paragraph_format.space_after = Pt(12)


def load_tex() -> str:
    return TEX_PATH.read_text(encoding="utf-8")


def extract_block(text: str, start_marker: str, end_marker: str) -> str:
    start = text.index(start_marker)
    end = text.index(end_marker, start)
    return text[start:end]


def strip_latex_commands(text: str, label_map: Dict[str, str]) -> str:
    text = text.replace(r"\_", "_").replace(r"\%", "%").replace(r"\ ", " ")
    text = re.sub(r"\\textsuperscript\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\texttt\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\textbf\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\textit\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\underline\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\url\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\text\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\eqref\{([^}]*)\}", lambda m: label_map.get(m.group(1), ""), text)
    text = re.sub(r"\\ref\{([^}]*)\}", lambda m: label_map.get(m.group(1), ""), text)
    text = re.sub(r"\\addcontentsline\{[^}]*\}\{[^}]*\}\{[^}]*\}", "", text)
    text = re.sub(r"\\zihao\{[^}]*\}", "", text)
    text = re.sub(r"\\heiti", "", text)
    text = re.sub(r"\\textbf", "", text)
    text = re.sub(r"\\noindent", "", text)
    text = text.replace(r"\quad", "  ")
    text = text.replace("~", " ")
    text = text.replace(r"\\", "\n")
    text = re.sub(r"\\[A-Za-z]+\*?(?:\[[^\]]*\])?(?:\{[^}]*\})?", "", text)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def equation_to_text(raw: str) -> str:
    text = raw.replace("\n", " ").strip()
    replacements = {
        r"\text{avg}": "avg",
        r"\text{p95}": "p95",
        r"\text{ ms}": " ms",
        r"\text{client}": "client",
        r"\text{Promise}": "Promise",
        r"\rightarrow": " -> ",
        r"\langle": "<",
        r"\rangle": ">",
        r"\leq": " <= ",
        r"\cdot": " * ",
        r"\quad": "  ",
        r"\sum_{i=1}^{n}": "Σ(i=1..n)",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace(r"\_", "_")
    text = re.sub(r"\\label\{[^}]*\}", "", text)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def build_label_map(text: str) -> Dict[str, str]:
    label_map: Dict[str, str] = {}
    current_chapter = 0
    figure_counts: Dict[int, int] = {}
    table_counts: Dict[int, int] = {}
    equation_counts: Dict[int, int] = {}
    env = None
    pending_kind = None

    for line in text.splitlines():
        if re.match(r"\\section\{", line):
            current_chapter += 1
        if r"\begin{figure}" in line:
            env = "figure"
        elif r"\begin{table}" in line:
            env = "table"
        elif r"\begin{equation}" in line:
            env = "equation"
            equation_counts[current_chapter] = equation_counts.get(current_chapter, 0) + 1
        elif r"\end{figure}" in line or r"\end{table}" in line or r"\end{equation}" in line:
            env = None

        if r"\caption{" in line:
            if env == "figure":
                figure_counts[current_chapter] = figure_counts.get(current_chapter, 0) + 1
                pending_kind = ("图", current_chapter, figure_counts[current_chapter])
            elif env == "table":
                table_counts[current_chapter] = table_counts.get(current_chapter, 0) + 1
                pending_kind = ("表", current_chapter, table_counts[current_chapter])

        label_match = re.search(r"\\label\{([^}]*)\}", line)
        if label_match:
            label = label_match.group(1)
            if env == "equation":
                num = equation_counts.get(current_chapter, 1)
                label_map[label] = f"({current_chapter}-{num})"
            elif pending_kind:
                kind, chapter, num = pending_kind
                label_map[label] = f"{kind}{chapter}-{num}"
                pending_kind = None

    return label_map


def add_cover(doc: Document) -> None:
    lines = [
        ("江西师大软件学院", 22, True),
        ("本科生毕业论文", 22, True),
        ("", 12, False),
        ("基于React与tRPC的网上书城系统", 18, True),
        ("设计与实现", 18, True),
        ("Design and Implementation of an Online Bookstore System", 14, False),
        ("Based on React and tRPC", 14, False),
    ]
    for text, size, bold in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12 if text else 6)
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体" if bold else "宋体")

    doc.add_paragraph("")
    info_rows = [
        ("学生姓名：", "________________________"),
        ("学    号：", "________________________"),
        ("所在学院：", "软件学院"),
        ("所学专业：", "软件工程"),
        ("指导教师：", "________________________"),
        ("完成时间：", "________________________"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for left, right in info_rows:
        row = table.add_row().cells
        row[0].text = left
        row[1].text = right
    doc.add_page_break()


def add_heading_paragraph(doc: Document, text: str, level: int, center: bool = False) -> None:
    size_map = {1: 16, 2: 14, 3: 12}
    before_map = {1: 12, 2: 10, 3: 8}
    after_map = {1: 8, 2: 6, 3: 4}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(before_map[level])
    p.paragraph_format.space_after = Pt(after_map[level])
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size_map[level])
    run.font.name = "Times New Roman"
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def add_abstracts(doc: Document, text: str, label_map: Dict[str, str]) -> None:
    cn_block = extract_block(text, "% ======== 中文摘要 ========", "% ======== 英文摘要 ========")
    en_block = extract_block(text, "% ======== 英文摘要 ========", "% ======== 目录 ========")

    def parse_abstract_block(block: str) -> Tuple[str, List[str], str]:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        title = "摘要"
        paragraphs: List[str] = []
        keywords = ""
        for line in lines:
            if line.startswith(r"\begin{center}") or line.startswith(r"\end{center}") or line.startswith("%"):
                continue
            if "摘要" in line and line.startswith("{"):
                title = "摘要" if "Abstract" not in line else "Abstract"
                continue
            if "关键词" in line or "Key words" in line:
                keywords = strip_latex_commands(line, label_map)
                continue
            if line.startswith(r"\vspace") or line.startswith(r"\newpage") or line.startswith(r"\pagenumbering") or line.startswith(r"\setcounter"):
                continue
            cleaned = strip_latex_commands(line, label_map)
            if cleaned:
                paragraphs.append(cleaned)
        return title, paragraphs, keywords

    for block in [cn_block, en_block]:
        title, paragraphs, keywords = parse_abstract_block(block)
        add_heading_paragraph(doc, title, 1, center=True)
        for para in paragraphs:
            doc.add_paragraph(para)
        if keywords:
            doc.add_paragraph(keywords)
        doc.add_page_break()


def collect_outline(text: str, label_map: Dict[str, str]) -> List[Tuple[int, str]]:
    outline: List[Tuple[int, str]] = []
    body = text.split("% ======== 第1章 绪论 ========")[1]
    body = body.split(r"\section*{致谢}")[0]
    for line in body.splitlines():
        stripped = line.strip()
        if stripped.startswith(r"\section{"):
            title = strip_latex_commands(re.search(r"\{(.+)\}", stripped).group(1), label_map)
            outline.append((1, title))
        elif stripped.startswith(r"\subsection{"):
            title = strip_latex_commands(re.search(r"\{(.+)\}", stripped).group(1), label_map)
            outline.append((2, title))
        elif stripped.startswith(r"\subsubsection{"):
            title = strip_latex_commands(re.search(r"\{(.+)\}", stripped).group(1), label_map)
            outline.append((3, title))
    return outline


def add_toc(doc: Document, outline: List[Tuple[int, str]]) -> None:
    add_heading_paragraph(doc, "目录", 1, center=True)
    for level, title in outline:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.7 * (level - 1))
        para.paragraph_format.space_after = Pt(2)
        para.add_run(title)
    doc.add_page_break()


def flush_paragraph(doc: Document, lines: List[str], label_map: Dict[str, str]) -> None:
    if not lines:
        return
    text = strip_latex_commands(" ".join(lines), label_map)
    if text:
        doc.add_paragraph(text)
    lines.clear()


def add_caption(doc: Document, kind: str, chapter: int, counter: int, caption: str) -> None:
    p = doc.add_paragraph(style="CaptionCn")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{kind} {chapter}-{counter} {caption}")


def insert_image(doc: Document, image_path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(image_path), width=Cm(15.8))


def add_table(doc: Document, rows: List[List[str]], chapter: int, counter: int, caption: str) -> None:
    add_caption(doc, "表", chapter, counter, caption)
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, cell in enumerate(rows[0]):
        hdr[i].text = cell
        for paragraph in hdr[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, cell in enumerate(row):
            cells[i].text = cell
    doc.add_paragraph("")


def parse_body(doc: Document, text: str, label_map: Dict[str, str]) -> None:
    body = text.split("% ======== 第1章 绪论 ========")[1]
    body = body.split(r"\end{document}")[0]
    lines = body.splitlines()

    current_chapter = 0
    figure_counter: Dict[int, int] = {}
    table_counter: Dict[int, int] = {}
    paragraph_lines: List[str] = []
    figure_data = None
    table_data = None
    equation_lines: List[str] = []

    for raw_line in lines:
        line = raw_line.strip()
        if not line or line.startswith("%"):
            if not figure_data and not table_data and not equation_lines:
                flush_paragraph(doc, paragraph_lines, label_map)
            continue

        if figure_data is not None:
            if line.startswith(r"\includegraphics"):
                image_rel = re.search(r"\{([^}]*)\}", line).group(1)
                figure_data["image"] = (THESIS_DIR / image_rel).resolve()
            elif line.startswith(r"\caption{"):
                figure_data["caption"] = strip_latex_commands(re.search(r"\{(.+)\}", line).group(1), label_map)
            elif line.startswith(r"\end{figure}"):
                figure_counter[current_chapter] = figure_counter.get(current_chapter, 0) + 1
                if figure_data.get("image"):
                    insert_image(doc, figure_data["image"])
                add_caption(doc, "图", current_chapter, figure_counter[current_chapter], figure_data.get("caption", ""))
                figure_data = None
            continue

        if table_data is not None:
            if line.startswith(r"\caption{"):
                table_data["caption"] = strip_latex_commands(re.search(r"\{(.+)\}", line).group(1), label_map)
            elif line.startswith(r"\begin{tabularx}") or line.startswith(r"\toprule") or line.startswith(r"\midrule") or line.startswith(r"\bottomrule"):
                continue
            elif line.startswith(r"\end{tabularx}"):
                continue
            elif line.startswith(r"\end{table}"):
                table_counter[current_chapter] = table_counter.get(current_chapter, 0) + 1
                add_table(doc, table_data["rows"], current_chapter, table_counter[current_chapter], table_data.get("caption", ""))
                table_data = None
            elif "&" in line and line.endswith(r"\\"):
                row = [strip_latex_commands(part.strip().rstrip("\\"), label_map) for part in line[:-2].split("&")]
                table_data["rows"].append(row)
            continue

        if equation_lines:
            if line.startswith(r"\end{equation}"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(equation_to_text(" ".join(equation_lines)))
                equation_lines = []
            else:
                equation_lines.append(line)
            continue

        if line.startswith(r"\section{"):
            flush_paragraph(doc, paragraph_lines, label_map)
            current_chapter += 1
            title = strip_latex_commands(re.search(r"\{(.+)\}", line).group(1), label_map)
            add_heading_paragraph(doc, title, 1)
            continue

        if line.startswith(r"\section*{"):
            flush_paragraph(doc, paragraph_lines, label_map)
            title = strip_latex_commands(re.search(r"\{(.+)\}", line).group(1), label_map)
            add_heading_paragraph(doc, title, 1)
            continue

        if line.startswith(r"\subsection{"):
            flush_paragraph(doc, paragraph_lines, label_map)
            title = strip_latex_commands(re.search(r"\{(.+)\}", line).group(1), label_map)
            add_heading_paragraph(doc, title, 2)
            continue

        if line.startswith(r"\subsubsection{"):
            flush_paragraph(doc, paragraph_lines, label_map)
            title = strip_latex_commands(re.search(r"\{(.+)\}", line).group(1), label_map)
            add_heading_paragraph(doc, title, 3)
            continue

        if line.startswith(r"\newpage"):
            flush_paragraph(doc, paragraph_lines, label_map)
            doc.add_page_break()
            continue

        if line.startswith(r"\begin{figure}"):
            flush_paragraph(doc, paragraph_lines, label_map)
            figure_data = {}
            continue

        if line.startswith(r"\begin{table}"):
            flush_paragraph(doc, paragraph_lines, label_map)
            table_data = {"rows": []}
            continue

        if line.startswith(r"\begin{equation}"):
            flush_paragraph(doc, paragraph_lines, label_map)
            equation_lines = []
            continue

        if line.startswith(r"\begin{enumerate}"):
            flush_paragraph(doc, paragraph_lines, label_map)
            continue

        if line.startswith(r"\end{enumerate}"):
            flush_paragraph(doc, paragraph_lines, label_map)
            continue

        if line.startswith(r"\item"):
            flush_paragraph(doc, paragraph_lines, label_map)
            item_text = strip_latex_commands(line[len(r"\item") :].strip(), label_map)
            doc.add_paragraph(item_text, style="List Number")
            continue

        if any(line.startswith(prefix) for prefix in [r"\pagenumbering", r"\setcounter", r"\tableofcontents", r"\addcontentsline"]):
            continue

        paragraph_lines.append(line)

    flush_paragraph(doc, paragraph_lines, label_map)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    tex = load_tex()
    label_map = build_label_map(tex)
    outline = collect_outline(tex, label_map)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_abstracts(doc, tex, label_map)
    add_toc(doc, outline)
    parse_body(doc, tex, label_map)
    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
